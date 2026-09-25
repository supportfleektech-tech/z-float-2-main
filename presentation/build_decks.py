#!/usr/bin/env python3
"""
Build the Z-float decks from one layout model.

  python3 presentation/build_decks.py

Outputs (next to this script):
  Z-float-System-Overview.pptx   10-slide system presentation
  Z-float-Pitch-Deck.pptx        investor / partner pitch deck
  html/overview.html, html/pitch.html   browser versions (same layout; used for visual QA)

Every element is positioned in inches on a 13.333 x 7.5 in (16:9) canvas and
rendered twice: python-pptx for PowerPoint/Keynote/Google Slides, and
absolute-positioned HTML (96 px/in) for the browser.
"""
from __future__ import annotations

import html
import os
from dataclasses import dataclass, field

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
W, H = 13.333, 7.5

NAVY = "0F2233"
BLUE = "1463FF"
SOFT = "EAF1FF"
INK = "0F172A"
MUTED = "5B6B83"
LINE = "D9E1EC"
GREEN = "0E9F6E"
AMBER = "B7791F"
RED = "C53030"
WHITE = "FFFFFF"
BG = "F7F9FC"
FONT = "Calibri"


# --------------------------------------------------------------------------- model
@dataclass
class Run:
    text: str
    bold: bool = False
    color: str | None = None


@dataclass
class Para:
    runs: list[Run]
    size: float = 16
    color: str = INK
    bold: bool = False
    align: str = "l"  # l | c | r
    space_after: float = 6
    bullet: bool = False


@dataclass
class Text:
    x: float
    y: float
    w: float
    h: float
    paras: list[Para]
    anchor: str = "t"  # t | m | b


@dataclass
class Rect:
    x: float
    y: float
    w: float
    h: float
    fill: str
    line: str | None = None
    radius: bool = True


@dataclass
class Image:
    x: float
    y: float
    w: float
    path: str
    h: float | None = None  # computed from aspect when None
    border: bool = True


@dataclass
class Slide:
    elements: list = field(default_factory=list)
    bg: str = WHITE
    notes: str = ""


def P(text: str | list, size=16, color=INK, bold=False, align="l", space_after=6, bullet=False) -> Para:
    runs = [Run(text)] if isinstance(text, str) else text
    return Para(runs, size, color, bold, align, space_after, bullet)


def img_size(path: str) -> tuple[int, int]:
    with open(path, "rb") as f:
        head = f.read(24)
    return int.from_bytes(head[16:20], "big"), int.from_bytes(head[20:24], "big")


def image(x, y, w, name, border=True, h=None) -> Image:
    path = os.path.join(ASSETS, name)
    pw, ph = img_size(path)
    return Image(x, y, w, path, h if h else w * ph / pw, border)


# --------------------------------------------------------------------------- building blocks
def header(slide: Slide, kicker: str, title: str, dark=False):
    slide.elements.append(Text(0.7, 0.45, 11.9, 0.35, [P(kicker.upper(), 12, "8FB3FF" if dark else BLUE, True)]))
    slide.elements.append(Text(0.7, 0.78, 11.9, 0.9, [P(title, 30, WHITE if dark else INK, True)]))


def footer(slide: Slide, n: int, total: int, dark=False):
    c = "8FA3BF" if dark else MUTED
    slide.elements.append(Text(0.7, 7.0, 6, 0.3, [P("Z-float · Business payments for Kenya", 10, c)]))
    slide.elements.append(Text(10.6, 7.0, 2.0, 0.3, [P(f"{n} / {total}", 10, c, align="r")]))


def card(slide: Slide, x, y, w, h, title, body: list[str], accent=BLUE, fill=WHITE, title_size=17, body_size=13):
    slide.elements.append(Rect(x, y, w, h, fill, LINE))
    slide.elements.append(Rect(x, y, 0.08, h, accent, None, radius=False))
    paras = [P(title, title_size, INK, True, space_after=8)]
    paras += [P(b, body_size, MUTED, bullet=True, space_after=4) for b in body]
    slide.elements.append(Text(x + 0.3, y + 0.2, w - 0.5, h - 0.3, paras))


def pill(slide: Slide, x, y, w, text, fill=SOFT, color=BLUE, size=12, h=0.38):
    slide.elements.append(Rect(x, y, w, h, fill, None))
    slide.elements.append(Text(x, y, w, h, [P(text, size, color, True, "c", 0)], anchor="m"))


def stat(slide: Slide, x, y, w, value, label, source="", dark=False):
    slide.elements.append(Text(x, y, w, 0.8, [P(value, 34, "8FB3FF" if dark else BLUE, True, space_after=0)]))
    slide.elements.append(Text(x, y + 0.8, w, 0.7, [P(label, 13, WHITE if dark else INK, space_after=2)]))
    if source:
        slide.elements.append(Text(x, y + 1.35, w, 0.4, [P(source, 9, "8FA3BF" if dark else MUTED)]))


def flow(slide: Slide, x, y, steps: list[tuple[str, str]], w=2.25, gap=0.35, h=1.25, fill=SOFT):
    for i, (t, s) in enumerate(steps):
        cx = x + i * (w + gap)
        slide.elements.append(Rect(cx, y, w, h, fill, None))
        slide.elements.append(Text(cx + 0.15, y + 0.12, w - 0.3, h - 0.2, [P(t, 14, INK, True, "c", 4), P(s, 11, MUTED, align="c", space_after=0)], anchor="m"))
        if i < len(steps) - 1:
            slide.elements.append(Text(cx + w, y + h / 2 - 0.25, gap, 0.5, [P("→", 20, BLUE, True, "c", 0)], anchor="m"))


def title_slide(title: str, subtitle: str, tag: str, notes: str) -> Slide:
    s = Slide(bg=NAVY, notes=notes)
    s.elements.append(Rect(0.7, 0.8, 0.62, 0.62, BLUE, None))
    s.elements.append(Text(0.7, 0.8, 0.62, 0.62, [P("Z", 26, WHITE, True, "c", 0)], anchor="m"))
    s.elements.append(Text(1.45, 0.86, 5, 0.55, [P("Z-float", 24, WHITE, True)], anchor="m"))
    s.elements.append(Text(0.7, 2.3, 8.2, 2.3, [P(title, 44, WHITE, True, space_after=0)]))
    s.elements.append(Text(0.7, 4.45, 7.6, 1.2, [P(subtitle, 18, "C9D6EA")]))
    s.elements.append(Text(0.7, 6.55, 8, 0.4, [P(tag, 12, "8FA3BF")]))
    s.elements.append(image(8.55, 1.5, 4.2, "receipt.png", border=False))
    return s


# --------------------------------------------------------------------------- 10-slide system overview
def overview() -> list[Slide]:
    T = 10
    slides: list[Slide] = []

    slides.append(title_slide(
        "Pay out. Get paid.\nStay KRA-compliant.",
        "Z-float is the business payments platform for Kenya — M-Pesa & bank payouts, inbound collections and KRA eTIMS invoicing on one double-entry ledger.",
        "System overview · 2026 release: collections, eTIMS e-invoicing & identity",
        "Introduce Z-float in one line: the operating account for a Kenyan business. This release adds the receiving side (collections), KRA eTIMS electronic invoices and receipts, and full identity (ID number + KRA PIN) on every person we deal with.",
    ))

    s = Slide(notes="Three pains we hear from finance teams: money moves in many rails, tax compliance is now digital and mandatory, and phone numbers alone do not identify people.")
    header(s, "The problem", "Kenyan businesses run money on scattered tools")
    card(s, 0.7, 1.9, 3.8, 3.9, "Money in, money out — separately", ["Payouts on one portal, M-Pesa paybill statements in another", "Manual matching of who paid which invoice", "No single balance the team trusts"], BLUE)
    card(s, 4.77, 1.9, 3.8, 3.9, "Tax compliance went digital", ["Every sale needs an eTIMS invoice or receipt", "Expenses without an eTIMS invoice are no longer deductible (Finance Act 2025)", "Penalty: double the tax due or KES 2M"], AMBER)
    card(s, 8.84, 1.9, 3.8, 3.9, "A phone number is not an identity", ["Numbers are shared, recycled, and mistyped", "Two “John Otieno”s in the payee book", "Buyers need their KRA PIN on invoices to claim VAT"], RED)
    s.elements.append(Text(0.7, 6.15, 11.9, 0.5, [P("Sources: Tax Procedures Act as amended by Finance Act 2025 (s.23A, eTIMS); KRA income & expense validation notice.", 10, MUTED)]))
    footer(s, 2, T); slides.append(s)

    s = Slide(bg=BG, notes="Four pillars. Everything posts to the same double-entry ledger, so the balance, the invoices and the tax records always agree.")
    header(s, "The platform", "One account for every shilling that moves")
    cols = [
        ("Pay out", BLUE, ["M-Pesa B2C, paybill, till & bank", "Bulk CSV, payroll, suppliers, bills, airtime", "Schedules & payment links"]),
        ("Get paid  · NEW", GREEN, ["M-Pesa STK request-to-pay", "Paybill/Till C2B with invoice routing", "Payment links & bank transfers"]),
        ("Comply  · NEW", AMBER, ["KRA eTIMS invoices, receipts, credit notes", "OSCU / VSCU device integration", "ID number + KRA PIN on every party"]),
        ("Control", NAVY, ["Maker-checker approvals & limits", "Double-entry ledger, reconciliation", "Audit trail, webhooks, public API"]),
    ]
    for i, (t, c, b) in enumerate(cols):
        card(s, 0.7 + i * 3.05, 1.9, 2.85, 3.3, t, b, c)
    flow(s, 0.7, 5.5, [("Customer pays", "STK · C2B · link · bank"), ("Wallet credited", "balanced journal"), ("Invoice settled", "PAID / PARTIAL"), ("eTIMS receipt", "KRA-signed, QR")], w=2.6, gap=0.5, h=1.1, fill=WHITE)
    footer(s, 3, T); slides.append(s)

    s = Slide(notes="The payout engine that already exists: every payment is idempotent, policy-checked and approved before it moves; the ledger reserves funds first.")
    header(s, "Paying out", "Disburse to anyone on M-Pesa or bank — safely")
    s.elements.append(Text(0.7, 1.85, 5.2, 4.8, [
        P("Single, bulk and scheduled payments to phones, tills, paybills and bank accounts.", 15, INK, space_after=12),
        P("Approval policies by amount, role and payee risk (maker-checker)", 14, MUTED, bullet=True),
        P("New payees need an independent approver before they receive funds", 14, MUTED, bullet=True),
        P("Funds reserved on the ledger before the provider is called", 14, MUTED, bullet=True),
        P("Idempotent APIs, provider circuit breakers, stuck-payment monitor", 14, MUTED, bullet=True),
        P("Payroll, supplier bills, airtime and expense claims built in", 14, MUTED, bullet=True),
    ]))
    s.elements.append(image(6.2, 1.8, 6.45, "dashboard.png"))
    footer(s, 4, T); slides.append(s)

    s = Slide(notes="New: the receiving side. Four channels, one collections ledger. Each successful payment credits the wallet through a balanced funding journal, settles the linked invoice and triggers an eTIMS receipt through the outbox.")
    header(s, "Receiving payments · new", "Collect by M-Pesa prompt, paybill, link or bank")
    s.elements.append(image(0.7, 1.8, 6.6, "collections.png"))
    s.elements.append(Text(7.65, 1.8, 5.0, 4.9, [
        P("M-Pesa STK push", 15, INK, True, space_after=2), P("PIN prompt on the payer’s phone; callback settles it.", 13, MUTED, space_after=10),
        P("Paybill / Till (C2B)", 15, INK, True, space_after=2), P("Account ACME-INV-000123 routes to the business and the invoice. Unknown accounts rejected at validation or parked for recon.", 13, MUTED, space_after=10),
        P("Payment links & bank transfers", 15, INK, True, space_after=2), P("Links cap uses atomically; bank references can only be recorded once.", 13, MUTED, space_after=10),
        P("Safe by construction", 15, INK, True, space_after=2), P("Idempotent on CheckoutRequestID / TransID; late callbacks after expiry still credit; row locks prevent double credit.", 13, MUTED),
    ]))
    footer(s, 5, T); slides.append(s)

    s = Slide(notes="eTIMS: we integrate with KRA's OSCU API (VSCU for high volume). Device initialisation returns a communication key we store encrypted. Each sale is sent to saveTrnsSalesOsdc; KRA returns the receipt number, signature, SCU ID and MRC which we print with the verification QR.")
    header(s, "KRA eTIMS · new", "Electronic tax invoices & receipts, signed by KRA")
    s.elements.append(image(0.7, 1.8, 4.7, "receipt.png"))
    s.elements.append(Text(5.7, 1.8, 7.0, 5.0, [
        P("What the business gets", 16, INK, True, space_after=6),
        P("Tax invoices, receipts and credit notes with VAT A–E computed per line", 13, MUTED, bullet=True),
        P("Automatic eTIMS receipt for every payment received", 13, MUTED, bullet=True),
        P("Buyer KRA PIN printed so customers can claim input VAT", 13, MUTED, bullet=True),
        P("Shareable link + QR that verifies on the KRA portal", 13, MUTED, bullet=True, space_after=14),
        P("How it works", 16, INK, True, space_after=6),
        P("OSCU API: selectInitOsdcInfo → cmcKey (encrypted at rest) → saveTrnsSalesOsdc", 13, MUTED, bullet=True),
        P("Sequential invoice numbers per device under a row lock", 13, MUTED, bullet=True),
        P("KRA outage? Documents queue and retry; rejections surface for a human", 13, MUTED, bullet=True),
        P("Signed documents are immutable in the database — corrections are credit notes", 13, MUTED, bullet=True),
    ]))
    footer(s, 6, T); slides.append(s)

    s = Slide(notes="Identity everywhere: recipients, customers, payers and team members carry ID type, ID number and KRA PIN, validated and normalised. One search box resolves a phone, ID number or KRA PIN.")
    header(s, "Identity · new", "ID number + KRA PIN — not just a phone number")
    s.elements.append(image(0.7, 1.8, 6.6, "recipients.png"))
    s.elements.append(Text(7.65, 1.8, 5.0, 4.9, [
        P("Captured on recipients, customers, payers and team members", 15, INK, True, space_after=10),
        P("National ID, Alien ID, Passport, Military ID, Company registration", 13, MUTED, bullet=True),
        P("KRA PIN validated (A/P + 9 digits + letter) and classified individual vs business", 13, MUTED, bullet=True),
        P("Duplicate guard: the same PIN or ID can’t be added twice by accident", 13, MUTED, bullet=True),
        P("One search box: phone, ID number or KRA PIN", 13, MUTED, bullet=True),
        P("ID numbers masked in lists; changes are audited", 13, MUTED, bullet=True),
        P("Payer PIN flows straight onto their eTIMS receipt", 13, MUTED, bullet=True),
    ]))
    footer(s, 7, T); slides.append(s)

    s = Slide(bg=BG, notes="Trust is the product. These are the invariants tests enforce on every change.")
    header(s, "Controls & trust", "Built like a bank ledger, not a spreadsheet")
    items = [
        ("Double-entry ledger", "Every credit and debit is a balanced journal; a health check proves it every 10 minutes."),
        ("Maker-checker", "Payments, reversals, new payees and privileged roles need an independent approver."),
        ("Idempotency everywhere", "API keys, provider callbacks (CheckoutRequestID, TransID) and bank refs dedupe."),
        ("Transactional outbox", "Receipts, notifications and webhooks are published exactly once, after commit."),
        ("Immutable fiscal records", "A database trigger blocks edits/deletes of KRA-signed documents."),
        ("Secrets & audit", "Device keys AES-GCM encrypted; callbacks token-checked; every change audited."),
    ]
    for i, (t, b) in enumerate(items):
        x = 0.7 + (i % 3) * 4.05
        y = 1.9 + (i // 3) * 2.35
        card(s, x, y, 3.85, 2.1, t, [b], [BLUE, GREEN, AMBER, NAVY, RED, BLUE][i], body_size=13)
    footer(s, 8, T); slides.append(s)

    s = Slide(notes="Monorepo: Next.js 14 web + API, a BullMQ worker and an outbox relay, all sharing typed packages. Providers and the eTIMS client are adapters, so sandbox and production differ only by configuration.")
    header(s, "Architecture", "Typed monorepo, adapters at the edges")
    layers = [
        ("Channels", "Portal (Next.js 14) · Public API · Hosted pay & receipt pages · Provider webhooks", SOFT),
        ("Domain packages", "payments-core · collections · etims · ledger · approvals · validation · auth/RBAC · audit", "E8F7F0"),
        ("Async", "Worker (BullMQ/Redis): executions, callbacks, reconciliation, eTIMS retries, STK expiry · Outbox relay", "FFF6E5"),
        ("Adapters", "M-Pesa Daraja (B2C, STK, C2B) · Bank PSP · Airtime · KRA eTIMS OSCU / VSCU / sandbox signer", "F3F4F8"),
        ("Data", "PostgreSQL (Drizzle, SQL migrations, triggers) · Redis · Object storage", "EEF2F7"),
    ]
    for i, (t, b, f) in enumerate(layers):
        y = 1.85 + i * 0.98
        s.elements.append(Rect(0.7, y, 11.9, 0.85, f, LINE))
        s.elements.append(Text(0.95, y, 2.6, 0.85, [P(t, 15, INK, True, space_after=0)], anchor="m"))
        s.elements.append(Text(3.6, y, 8.8, 0.85, [P(b, 13, MUTED, space_after=0)], anchor="m"))
    footer(s, 9, T); slides.append(s)

    s = Slide(bg=NAVY, notes="Close with what it takes to go live: KRA OSCU approval and device serial, Safaricom Daraja production credentials with STK and C2B, and the configuration flags. Then the roadmap.")
    header(s, "Go-live & roadmap", "From sandbox to production", dark=True)
    card(s, 0.7, 1.9, 5.8, 4.6, "Go-live checklist", [
        "KRA: apply for OSCU (or VSCU) on the eTIMS portal / GavaConnect; receive device serial",
        "Set ETIMS_DRIVER=oscu, ETIMS_ENVIRONMENT=production",
        "Map items to KRA item classification codes (placeholder today)",
        "Safaricom Daraja production app: STK passkey, C2B URL registration",
        "Set MPESA_C2B_CALLBACK_TOKEN; enable external validation",
        "Run the certification test cases in the KRA sandbox first",
    ], GREEN, fill="16304A", title_size=17, body_size=13)
    card(s, 6.84, 1.9, 5.8, 4.6, "Roadmap", [
        "Item catalogue synced with KRA (saveItem, stock movements)",
        "Purchase invoices: pull supplier eTIMS invoices to validate expenses",
        "Recurring invoices & automatic payment reminders by SMS",
        "Customer statements and aged-receivables report",
        "Card and Airtel Money collections",
        "Multi-branch eTIMS devices (bhfId per branch)",
    ], BLUE, fill="16304A", title_size=17, body_size=13)
    for e in s.elements:
        if isinstance(e, Text):
            for p in e.paras:
                if p.color == INK:
                    p.color = WHITE
                if p.color == MUTED:
                    p.color = "C9D6EA"
    footer(s, 10, T, dark=True); slides.append(s)
    return slides


# --------------------------------------------------------------------------- pitch deck
def pitch() -> list[Slide]:
    T = 12
    slides: list[Slide] = []
    slides.append(title_slide(
        "The money & tax OS\nfor Kenyan businesses",
        "Pay suppliers and staff, collect from customers, and issue KRA eTIMS invoices — from one account.",
        "Pitch deck · 2026",
        "Hook: every Kenyan business must now invoice through KRA eTIMS, and nearly all of them already move money on M-Pesa. Z-float joins the two.",
    ))

    s = Slide(notes="Two forces collide: mobile money is how Kenya pays, and eTIMS is now how Kenya taxes.")
    header(s, "Why now", "M-Pesa is how Kenya pays. eTIMS is how it taxes.")
    stat(s, 0.7, 2.0, 3.8, "KES 38.3T", "moved on M-Pesa in FY2025 (37.2 billion transactions)", "Safaricom Annual Report 2025")
    stat(s, 4.77, 2.0, 3.8, "2.4M", "M-Pesa merchants — 0.9M Lipa na M-Pesa tills + 1.5M Pochi la Biashara", "Safaricom H1 FY2026 results (Nov 2025)")
    stat(s, 8.84, 2.0, 3.8, "0%", "of an expense is deductible without an eTIMS invoice", "Tax Procedures Act, Finance Act 2025")
    s.elements.append(Rect(0.7, 4.35, 11.9, 1.9, SOFT, None))
    s.elements.append(Text(1.0, 4.5, 11.3, 1.6, [
        P("From the 2025 accounting period KRA validates declared income and expenses against eTIMS data on iTax. Failing to issue an eTIMS invoice costs double the tax due or KES 2 million, whichever is higher.", 16, INK, space_after=6),
        P("Every business that collects on M-Pesa now needs a compliant invoice or receipt for each sale — and a way to link it to the payment.", 16, BLUE, True),
    ], anchor="m"))
    footer(s, 2, T); slides.append(s)

    s = Slide(notes="Today a finance team stitches together 4–5 tools and a spreadsheet.")
    header(s, "Problem", "Five tools and a spreadsheet to run the money")
    flow(s, 0.7, 2.1, [("M-Pesa portal", "payouts one by one"), ("Paybill statement", "who paid what?"), ("Invoicing app", "not linked to payments"), ("eTIMS portal", "retype every sale"), ("Spreadsheet", "reconcile by hand")], w=2.05, gap=0.41, h=1.5, fill="FDECEC")
    s.elements.append(Text(0.7, 4.1, 11.9, 2.4, [
        P("Hours lost every week matching payments to invoices", 16, INK, bullet=True),
        P("Tax exposure: sales without eTIMS invoices, expenses disallowed", 16, INK, bullet=True),
        P("Fraud risk: payees identified only by phone numbers, no approvals", 16, INK, bullet=True),
        P("No single balance that finance, owners and auditors agree on", 16, INK, bullet=True),
    ]))
    footer(s, 3, T); slides.append(s)

    s = Slide(bg=BG, notes="Z-float replaces the patchwork with one account and one ledger.")
    header(s, "Solution", "One account: pay out, get paid, stay compliant")
    for i, (t, c, b) in enumerate([
        ("Pay out", BLUE, ["Bulk & single M-Pesa / bank payouts", "Payroll, suppliers, bills", "Approvals & limits"]),
        ("Get paid", GREEN, ["STK request-to-pay", "Paybill routing to invoices", "Links & bank transfers"]),
        ("Comply", AMBER, ["KRA eTIMS invoices & receipts", "Auto receipt on every payment", "ID + KRA PIN on every party"]),
    ]):
        card(s, 0.7 + i * 4.05, 1.9, 3.85, 2.6, t, b, c, title_size=20, body_size=14)
    s.elements.append(Text(0.7, 4.85, 11.9, 1.6, [
        P("The loop no one else closes:", 18, INK, True, "c", 6),
        P("invoice  →  M-Pesa prompt  →  wallet credited  →  invoice PAID  →  KRA-signed receipt", 20, BLUE, True, "c"),
    ]))
    footer(s, 4, T); slides.append(s)

    s = Slide(notes="Product walkthrough: issue a tax invoice with VAT computed per line, send an STK prompt, and the receipt is issued automatically.")
    header(s, "Product", "Invoice → prompt → paid → KRA receipt, in one screen")
    s.elements.append(image(0.7, 1.8, 5.95, "newinvoice.png"))
    s.elements.append(image(6.85, 1.8, 5.8, "invoices.png"))
    s.elements.append(Text(0.7, 6.05, 11.9, 0.6, [P("Live product screens (sandbox). Left: new tax invoice with KRA tax bands. Right: signed invoices, receipts and credit notes with KRA CU numbers.", 12, MUTED)]))
    footer(s, 5, T); slides.append(s)

    s = Slide(notes="Collections view: every channel lands in one list, with payer identity and the eTIMS receipt.")
    header(s, "Product", "Every shilling in, with who paid and the receipt")
    s.elements.append(image(0.7, 1.8, 7.2, "collections.png"))
    s.elements.append(Text(8.25, 1.9, 4.4, 4.8, [
        P("Received today, 30-day totals, pending prompts", 14, MUTED, bullet=True),
        P("Payer name, phone, ID and KRA PIN", 14, MUTED, bullet=True),
        P("M-Pesa receipt number + linked eTIMS receipt", 14, MUTED, bullet=True),
        P("Unmatched paybill money parked for reconciliation — never lost", 14, MUTED, bullet=True),
        P("Webhooks: collection.received, invoice.fiscalised", 14, MUTED, bullet=True),
    ]))
    footer(s, 6, T); slides.append(s)

    s = Slide(notes="Who we sell to first.")
    header(s, "Customers", "Built for businesses that already live on M-Pesa")
    for i, (t, b) in enumerate([
        ("Distributors & wholesalers", "High volumes of paybill receipts to match to invoices; VAT-registered buyers demand PIN-bearing invoices."),
        ("Service businesses", "Clinics, schools, agencies: request-to-pay and automatic receipts replace manual books."),
        ("Employers & agri-aggregators", "Bulk payouts to staff and farmers, with ID-verified payees and approvals."),
        ("Accountants & bookkeepers", "Manage many clients' eTIMS compliance and reconciliation from one login."),
    ]):
        card(s, 0.7 + (i % 2) * 6.05, 1.9 + (i // 2) * 2.35, 5.85, 2.1, t, [b], [BLUE, GREEN, AMBER, NAVY][i], title_size=18, body_size=14)
    footer(s, 7, T); slides.append(s)

    s = Slide(bg=BG, notes="Revenue lines. Numbers are illustrative pricing to validate with design partners.")
    header(s, "Business model", "Subscription + transaction revenue")
    for i, (t, v, b) in enumerate([
        ("Platform subscription", "per month", ["Tiered by users & eTIMS devices", "Includes invoices, receipts, approvals"]),
        ("Payout fees", "per transaction", ["Margin over M-Pesa / bank tariffs", "Bulk & payroll bundles"]),
        ("Collection fees", "% of value", ["STK, paybill routing, payment links", "Waived on higher tiers"]),
        ("Compliance add-ons", "per client", ["Accountant multi-client workspace", "Expense validation & reports"]),
    ]):
        x = 0.7 + i * 3.05
        s.elements.append(Rect(x, 1.9, 2.85, 3.6, WHITE, LINE))
        s.elements.append(Text(x + 0.25, 2.1, 2.4, 3.3, [P(t, 17, INK, True, space_after=2), P(v, 13, BLUE, True, space_after=12)] + [P(b_, 13, MUTED, bullet=True) for b_ in b]))
    s.elements.append(Text(0.7, 5.8, 11.9, 0.8, [P("Pricing to be validated with design partners; fee levels are set per tier in the platform’s fee schedules today.", 12, MUTED)]))
    footer(s, 8, T); slides.append(s)

    s = Slide(notes="Where we compete and why we win.")
    header(s, "Competition", "Others do one piece. Z-float closes the loop.")
    rows = [
        ("", "Payouts", "Collections", "eTIMS invoices", "Invoice ↔ payment link", "Approvals & ledger"),
        ("M-Pesa business portals", "Yes", "Yes", "No", "No", "No"),
        ("Invoicing / accounting apps", "No", "Partial", "Yes", "Partial", "No"),
        ("Payment gateways", "Partial", "Yes", "No", "Partial", "No"),
        ("Z-float", "Yes", "Yes", "Yes", "Yes", "Yes"),
    ]
    cw = [3.4, 1.7, 1.7, 1.7, 1.7, 1.7]
    for r, row in enumerate(rows):
        y = 1.95 + r * 0.8
        fill = NAVY if r == 0 else (SOFT if row[0] == "Z-float" else WHITE)
        s.elements.append(Rect(0.7, y, 11.9, 0.72, fill, LINE, radius=False))
        x = 0.7
        for c, val in enumerate(row):
            color = WHITE if r == 0 else (GREEN if val == "Yes" else (MUTED if val in ("No", "Partial") else INK))
            s.elements.append(Text(x + 0.15, y, cw[c] - 0.2, 0.72, [P(val, 14 if r else 13, color, r == 0 or row[0] == "Z-float" or val == "Yes", "l" if c == 0 else "c", 0)], anchor="m"))
            x += cw[c]
    s.elements.append(Text(0.7, 6.1, 11.9, 0.5, [P("Category comparison, not a claim about any named vendor.", 11, MUTED)]))
    footer(s, 9, T); slides.append(s)

    s = Slide(notes="What is real today.")
    header(s, "Traction & readiness", "Production-grade engineering, ready for pilots")
    stat(s, 0.7, 1.95, 2.9, "4", "collection channels live in sandbox")
    stat(s, 3.75, 1.95, 2.9, "5", "KRA tax bands (A–E) computed per line")
    stat(s, 6.8, 1.95, 2.9, "100%", "of payments balanced on a double-entry ledger")
    stat(s, 9.85, 1.95, 2.9, "0", "editable fields on a signed KRA document")
    s.elements.append(Text(0.7, 4.1, 11.9, 2.5, [
        P("Built: payouts, approvals, reconciliation, collections, eTIMS (OSCU/VSCU adapter + sandbox signer), identity", 15, INK, bullet=True),
        P("Automated tests on real PostgreSQL cover money movement, idempotency and fiscal immutability", 15, INK, bullet=True),
        P("Next: KRA OSCU certification + Safaricom production go-live with design partners", 15, INK, bullet=True),
    ]))
    footer(s, 10, T); slides.append(s)

    s = Slide(bg=BG, notes="Milestones for the next 12 months.")
    header(s, "Roadmap", "Next 12 months")
    flow(s, 0.7, 2.0, [("Q1", "KRA OSCU certification · Daraja production · 10 design partners"), ("Q2", "Item catalogue sync · recurring invoices · SMS reminders"), ("Q3", "Purchase-invoice validation · accountant workspace"), ("Q4", "Airtel Money & cards · multi-branch devices")], w=2.7, gap=0.37, h=2.2, fill=WHITE)
    s.elements.append(Text(0.7, 4.7, 11.9, 1.5, [P("Goal: the default way a Kenyan SME moves money and proves it to KRA.", 20, BLUE, True, "c")], anchor="m"))
    footer(s, 11, T); slides.append(s)

    s = Slide(bg=NAVY, notes="Close: what we are asking for.")
    header(s, "The ask", "Partner with us to go live", dark=True)
    card(s, 0.7, 1.9, 5.8, 4.3, "We are looking for", [
        "Design-partner businesses (distributors, clinics, schools)",
        "Accounting firms managing eTIMS for SMEs",
        "Banking / PSP partners for settlement and collections",
        "Seed funding for certification, go-to-market and support",
    ], GREEN, fill="16304A", title_size=18, body_size=14)
    card(s, 6.84, 1.9, 5.8, 4.3, "What you get", [
        "Early access and influence on the roadmap",
        "Preferential pricing for the first year",
        "Hands-on onboarding to eTIMS and M-Pesa collections",
        "A single view of money in, money out and tax",
    ], BLUE, fill="16304A", title_size=18, body_size=14)
    for e in s.elements:
        if isinstance(e, Text):
            for p in e.paras:
                if p.color == INK:
                    p.color = WHITE
                if p.color == MUTED:
                    p.color = "C9D6EA"
    footer(s, 12, T, dark=True); slides.append(s)
    return slides


# --------------------------------------------------------------------------- renderers
def hexrgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.upper())


def to_pptx(slides: list[Slide], path: str):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    for sd in slides:
        sl = prs.slides.add_slide(blank)
        bg = sl.background.fill
        bg.solid()
        bg.fore_color.rgb = hexrgb(sd.bg)
        for e in sd.elements:
            if isinstance(e, Rect):
                shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if e.radius else MSO_SHAPE.RECTANGLE, Inches(e.x), Inches(e.y), Inches(e.w), Inches(e.h))
                if e.radius:
                    shp.adjustments[0] = min(0.12, 0.12 / max(e.h, 0.1))
                shp.fill.solid()
                shp.fill.fore_color.rgb = hexrgb(e.fill)
                if e.line:
                    shp.line.color.rgb = hexrgb(e.line)
                    shp.line.width = Pt(0.75)
                else:
                    shp.line.fill.background()
                shp.shadow.inherit = False
                shp.text_frame.text = ""
            elif isinstance(e, Image):
                pic = sl.shapes.add_picture(e.path, Inches(e.x), Inches(e.y), Inches(e.w), Inches(e.h))
                if e.h is not None:
                    # crop to box from the top (keep aspect; show the top of the screenshot)
                    pw, ph = img_size(e.path)
                    natural_h = e.w * ph / pw
                    if natural_h > e.h + 0.01:
                        pic.crop_bottom = 1 - e.h / natural_h
                if e.border:
                    pic.line.color.rgb = hexrgb(LINE)
                    pic.line.width = Pt(0.75)
            elif isinstance(e, Text):
                tb = sl.shapes.add_textbox(Inches(e.x), Inches(e.y), Inches(e.w), Inches(e.h))
                tf = tb.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
                tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[e.anchor]
                for i, p in enumerate(e.paras):
                    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[p.align]
                    para.space_after = Pt(p.space_after)
                    lines = []
                    for r in p.runs:
                        lines.append(r)
                    if p.bullet:
                        lines = [Run("•  ", False, BLUE)] + lines
                    for r in lines:
                        chunks = r.text.split("\n")
                        for ci, chunk in enumerate(chunks):
                            if ci > 0:
                                para.add_line_break()
                            run = para.add_run()
                            run.text = chunk
                            f = run.font
                            f.name = FONT
                            f.size = Pt(p.size)
                            f.bold = r.bold or p.bold
                            f.color.rgb = hexrgb(r.color or p.color)
        if sd.notes:
            sl.notes_slide.notes_text_frame.text = sd.notes
    prs.save(path)


def to_html(slides: list[Slide], path: str, title: str):
    px = 96
    out = [f"""<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(title)}</title>
<style>
body{{margin:0;background:#1f2937;font-family:Calibri,Carlito,'Segoe UI',Inter,Arial,sans-serif}}
.slide{{position:relative;width:{W*px:.0f}px;height:{H*px:.0f}px;margin:24px auto;overflow:hidden;box-shadow:0 8px 30px rgba(0,0,0,.35)}}
.el{{position:absolute;box-sizing:border-box}}
.t{{display:flex;flex-direction:column;white-space:normal;line-height:1.18}}
.t p{{margin:0}}
img.el{{object-fit:cover;object-position:top}}
@media print{{body{{background:#fff}}.slide{{margin:0;box-shadow:none;page-break-after:always}}}}
</style></head><body>"""]
    rel = os.path.relpath(ASSETS, os.path.dirname(path))
    for sd in slides:
        out.append(f'<section class="slide" style="background:#{sd.bg}">')
        for e in sd.elements:
            box = f"left:{e.x*px:.1f}px;top:{e.y*px:.1f}px;width:{e.w*px:.1f}px;"
            if isinstance(e, Rect):
                r = "8px" if e.radius else "0"
                border = f"border:1px solid #{e.line};" if e.line else ""
                out.append(f'<div class="el" style="{box}height:{e.h*px:.1f}px;background:#{e.fill};{border}border-radius:{r}"></div>')
            elif isinstance(e, Image):
                border = f"border:1px solid #{LINE};" if e.border else ""
                src = os.path.join(rel, os.path.basename(e.path))
                out.append(f'<img class="el" src="{src}" style="{box}height:{e.h*px:.1f}px;{border}">')
            elif isinstance(e, Text):
                just = {"t": "flex-start", "m": "center", "b": "flex-end"}[e.anchor]
                out.append(f'<div class="el t" style="{box}height:{e.h*px:.1f}px;justify-content:{just}">')
                for p in e.paras:
                    align = {"l": "left", "c": "center", "r": "right"}[p.align]
                    runs = ""
                    if p.bullet:
                        runs += f'<span style="color:#{BLUE}">•&nbsp;&nbsp;</span>'
                    for r in p.runs:
                        t = html.escape(r.text).replace("\n", "<br>")
                        runs += f'<span style="color:#{r.color or p.color};font-weight:{700 if (r.bold or p.bold) else 400}">{t}</span>'
                    out.append(f'<p style="font-size:{p.size*px/72:.1f}px;text-align:{align};margin-bottom:{p.space_after*px/72:.1f}px">{runs}</p>')
                out.append("</div>")
        out.append("</section>")
    out.append("</body></html>")
    with open(path, "w") as f:
        f.write("\n".join(out))


def main():
    os.makedirs(os.path.join(HERE, "html"), exist_ok=True)
    ov, pt = overview(), pitch()
    assert len(ov) == 10, len(ov)
    to_pptx(ov, os.path.join(HERE, "Z-float-System-Overview.pptx"))
    to_pptx(pt, os.path.join(HERE, "Z-float-Pitch-Deck.pptx"))
    to_html(ov, os.path.join(HERE, "html", "overview.html"), "Z-float — System overview")
    to_html(pt, os.path.join(HERE, "html", "pitch.html"), "Z-float — Pitch deck")
    print(f"overview: {len(ov)} slides, pitch: {len(pt)} slides")


if __name__ == "__main__":
    main()
